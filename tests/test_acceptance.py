"""Synthetic, owner-authored rule expectations. Not an external gold standard."""
import copy
import csv
import io
from datetime import date, datetime
import pytest
from openpyxl import load_workbook, Workbook
from corroborai.demo import make_files
from corroborai.engine import load_inputs, reconcile, contract_code, history_bounds, match_assignments, expected_value
from corroborai.io import InputError, day, table, csv_bytes, xlsx_bytes, boolean, number


def changed(files, kind, column, value, row=2):
    files = dict(files)
    name, raw = files[kind]
    wb = load_workbook(io.BytesIO(raw)); ws = wb.active
    headers = [c.value for c in ws[1]]
    ws.cell(row, headers.index(column) + 1).value = value
    out = io.BytesIO(); wb.save(out); wb.close(); files[kind] = (name, out.getvalue())
    return files


def check(run, field, person="1001"):
    return next(r for r in run["results"] if r["field"] == field and r["person_id"] == person)


@pytest.fixture
def tables():
    return load_inputs(make_files())[0]


def test_full_demo_independently_authored_expected_cases():
    run = reconcile(make_files())
    assert run["metadata"]["assignment_pairs"] == 4
    assert run["metadata"]["source_rows"] == 5
    assert run["metadata"]["destination_rows"] == 4
    assert len(run["results"]) == 101
    assert check(run, "givenName")["verdict"] == "compliant"
    assert check(run, "contactEmail")["verdict"] == "justified_difference"
    assert check(run, "contactEmail")["expected"] == "cexemple001@loto-quebec.com"
    assert check(run, "assignmentStartDate")["expected"] == "2024-01-01"
    assert check(run, "weeklyHoursOverride", "1002")["expected"] == 35
    assert check(run, "weeklyHoursOverride", "1002")["verdict"] == "anomaly"
    assert check(run, "statusReasonCode", "1003")["expected"] == "EXTABS"
    assert check(run, "detailedStatus", "1003")["expected"] == "Absence complète"
    assert check(run, "__assignment__", "1004")["verdict"] == "anomaly"
    assert check(run, "positionName", "1002")["suggestions"]["learned"]["candidates"][0]["label"] == "502-Services financiers"
    assert len(run["metadata"]["mapped_fields"]) == 25
    assert "activityStatus" not in run["metadata"]["mapped_fields"]


@pytest.mark.parametrize("category,permanent,full,expected", [("V",1,1,"JWN"),("V",1,0,"XFLR"),("T",0,0,"KELH"),("O",0,0,"WHX"),("M",0,0,"CEGQ"),("R",0,0,"CNZC"),("J",0,0,"RMQ"),("Z",0,0,"JAW"),("Q",0,0,"TRSY")])
def test_contract_branches(category, permanent, full, expected):
    assert contract_code({"CatégorieEmploi":category,"EstPermanent":permanent,"EstTempsPlein":full}) == expected


@pytest.mark.parametrize("category,permanent,full", [("V",0,1),("V",1,"unknown"),("X",1,1)])
def test_uncovered_contract_abstains(category, permanent, full):
    with pytest.raises(InputError):
        contract_code({"CatégorieEmploi":category,"EstPermanent":permanent,"EstTempsPlein":full})


@pytest.mark.parametrize("access,status,reason", [("00","Actif",None),("01","Actif",None),("02","Absence complète","EXTABS"),("03","Absence complète","EXTABS"),("06","Absence complète","EXTABS"),("07","Absence complète","EXTABS")])
def test_every_access_branch(tables, access, status, reason):
    row = copy.deepcopy(tables["source"][2]); row["CodeSuspensionAccès"] = access
    tables["reasons"][0]["CodeGestionAccès"] = access
    assert expected_value(row,"detailedStatus",tables)[0] == status
    assert expected_value(row,"statusReasonCode",tables)[0] == reason
    returned = expected_value(row,"expectedReturnDate",tables)[0]
    assert returned == (None if access in ("00","01") else date(2025,2,1))


def test_absence_missing_or_duplicate_lookup_review(tables):
    row=tables["source"][2]
    tables["reasons"] = []
    with pytest.raises(InputError): expected_value(row,"statusReasonCode",tables)
    assert expected_value(tables["source"][0],"statusReasonCode",tables)[0] is None


def test_lookup_conflict_and_unknown_access(tables):
    row=tables["source"][2]; tables["reasons"][0]["CodeGestionAccès"]="07"
    with pytest.raises(InputError): expected_value(row,"statusReasonCode",tables)
    row["CodeSuspensionAccès"]="99"
    with pytest.raises(InputError): expected_value(row,"detailedStatus",tables)


@pytest.mark.parametrize("typ,primary,temporary",[("P",True,False),("A",False,True),("S",False,False)])
def test_assignment_flags(tables,typ,primary,temporary):
    row=tables["source"][0];row["TypeAffectation"]=typ
    assert expected_value(row,"isPrimaryAssignment",tables)[0] is primary
    assert expected_value(row,"isTemporaryAssignment",tables)[0] is temporary


@pytest.mark.parametrize("email,verdict",[("cexemple001@loto-quebec.com","justified_difference"),("dev-08-v2_cexemple001@loto-quebec.com","justified_difference"),("other_cexemple001@loto-quebec.com","review_required"),("cexemple002@loto-quebec.com","anomaly"),("cexemple001@evil.invalid","anomaly")])
def test_email_exact_base_and_prefixes(email,verdict):
    run=reconcile(changed(make_files(),"destination","contactEmail",email))
    assert check(run,"contactEmail")["verdict"] == verdict


def test_declared_new_prefix_and_accent_removal():
    files=changed(make_files(),"destination","contactEmail","stage_cexemple001@loto-quebec.com")
    assert check(reconcile(files,prefixes=["stage_"]),"contactEmail")["verdict"]=="justified_difference"
    files=changed(make_files(),"source","PrénomUsuel","Cámille")
    assert check(reconcile(files),"contactEmail")["expected"]=="cexemple001@loto-quebec.com"
    assert check(reconcile(files),"givenName")["verdict"]=="anomaly"


def position(tables, when, unit, rownum):
    row=copy.deepcopy(tables["positions"][0]);row.update(DateEffetAffectation=when,CodeDirectionAffectée=unit,_row=rownum);return row


def test_history_repeated_unit_and_next_change(tables):
    rows=[position(tables,datetime(2024,1,1),"10",2),position(tables,datetime(2024,2,1),"10",3),position(tables,datetime(2025,1,1),"20",4)]
    start,end,proof=history_bounds(tables["source"][0],rows)
    assert start==date(2024,1,1) and end==date(2024,12,31)
    assert {r["row"] for r in proof}=={2,3,4}
    tables["source"][0]["DateSortiePoste"]=datetime(2024,7,1)
    assert history_bounds(tables["source"][0],rows)[1]==date(2024,7,1)


def test_history_no_change_no_extra_end_and_minimum_start(tables):
    row=tables["source"][0]
    assert history_bounds(row,tables["positions"])[1] is None
    row["DateEntréePoste"]=datetime(2023,1,1)
    assert history_bounds(row,tables["positions"])[0]==date(2023,1,1)


def test_history_return_to_unit_requires_reference(tables):
    rows=[position(tables,datetime(2024,1,1),"10",2),position(tables,datetime(2024,4,1),"20",3),position(tables,datetime(2024,7,1),"10",4)]
    with pytest.raises(InputError):history_bounds(tables["source"][0],rows)
    assert history_bounds(tables["source"][0],rows,date(2024,2,1))[1]==date(2024,3,31)


def test_history_duplicates_missing_and_contradictions(tables):
    row=tables["source"][0];hist=[position(tables,datetime(2024,1,1),"10",2)]
    with pytest.raises(InputError):history_bounds(row,[])
    with pytest.raises(InputError):history_bounds(row,hist+hist)
    row["DateSortiePoste"]=datetime(2020,1,1)
    with pytest.raises(InputError):history_bounds(row,hist)


def test_matching_never_reuses_rows_and_retains_ambiguity(tables):
    src=tables["source"];dst=tables["destination"]
    pairs,left=match_assignments(src,dst)
    assert len({id(b) for _,b,_ in pairs})==len(pairs)==4 and len(left)==1
    a=copy.deepcopy(src[0]);a["_row"]=100
    pairs,left=match_assignments([src[0],a],[dst[0]])
    assert not pairs and len(left)==3 and all(r[2]=="review_required" for r in left)


def test_matching_same_role_two_types_and_wrong_flag(tables):
    a=copy.deepcopy(tables["source"][0]);a["TypeAffectation"]="A";a["_row"]=100
    pairs,left=match_assignments([tables["source"][0],a],[tables["destination"][0]])
    assert len(pairs)==1 and pairs[0][0]["TypeAffectation"]=="P" and len(left)==1
    files=changed(make_files(),"destination","isTemporaryAssignment",True)
    assert check(reconcile(files),"isTemporaryAssignment")["verdict"]=="anomaly"


def test_missing_person_and_role_mismatch_do_not_hide_error(tables):
    source=tables["source"][0];dest=tables["destination"][0]
    dest["positionId"]="WRONG"
    pairs,left=match_assignments([source],[dest]);assert len(pairs)==1 and pairs[0][2].startswith("inferred_")
    source["Matricule"]="";pairs,left=match_assignments([source],[dest]);assert len(left)==2


def test_input_errors_and_mapping_fail_closed():
    files=make_files();bad=dict(files);bad["source"]=("bad.xlsx",b"not a workbook")
    with pytest.raises(InputError):reconcile(bad)
    name,raw=files["source"];wb=load_workbook(io.BytesIO(raw));wb.active.cell(1,1).value="wrongHeader";buf=io.BytesIO();wb.save(buf);bad["source"]=(name,buf.getvalue())
    with pytest.raises(InputError):reconcile(bad)
    bad=changed(files,"mapping","Rule","Different unreviewed rule")
    with pytest.raises(InputError,match="unreviewed"):reconcile(bad)
    with pytest.raises(InputError):reconcile({})
    with pytest.raises(InputError):reconcile(changed(files,"source","Matricule","=1+1"))


@pytest.mark.parametrize("kind,sheet",[("source","Employe_Source"),("destination","Employe_Destination"),("positions","Feuil1"),("reasons","Sheet1")])
def test_blank_required_sheet_is_actionable_input_error(kind,sheet):
    files=make_files();wb=Workbook();wb.active.title=sheet;wb.active['A1']=''
    buf=io.BytesIO();wb.save(buf);wb.close();files[kind]=('empty.xlsx',buf.getvalue())
    with pytest.raises(InputError,match='empty worksheet|missing header'):
        reconcile(files)


def test_embedded_csv_shape_and_excel_dates(tables):
    assert len(tables["positions"])==10
    assert day("45292",datetime(1899,12,30))==date(2024,1,1)
    assert day("2024-01-01",datetime(1899,12,30))==date(2024,1,1)
    with pytest.raises(InputError):day("01/02/2024",datetime(1899,12,30))
    with pytest.raises(InputError):day("45292.5",datetime(1899,12,30))
    with pytest.raises(InputError):boolean("yes")
    assert boolean("Oui") is True and boolean("Non") is False
    with pytest.raises(InputError):number("NaN")


def test_exports_keep_evidence_and_prevent_formulas():
    run=reconcile(make_files());run["results"][0]["source_value"]="=HYPERLINK(\"evil\")"
    rows=list(csv.DictReader(io.StringIO(csv_bytes(run).decode('utf-8-sig'))))
    assert len(rows)==101 and rows[0]["source_value"].startswith("'=")
    assert all(r['reason'] and r['rule_id'] and (r['source_ref']!='null' or r['target_ref']!='null') for r in rows)
    wb=load_workbook(io.BytesIO(xlsx_bytes(run)),data_only=False)
    assert set(wb.sheetnames)=={"Audit","Investigation","Justified","Run metadata"}
    assert wb['Audit'].max_row==102
    assert wb['Audit'].cell(2,6).data_type=='s'
    assert not any(c.data_type=='f' for ws in wb for row in ws for c in row)


def test_ai_has_no_power_to_change_verdicts_or_identity():
    a=reconcile(make_files());b=reconcile(make_files(),use_ai=False)
    assert [(r['id'],r['verdict'],r['expected']) for r in a['results']]==[(r['id'],r['verdict'],r['expected']) for r in b['results']]
    assert all(not r['suggestions'] for r in a['results'] if r['field'] in ['personId','givenName','surname','contactEmail','detailedStatus'])

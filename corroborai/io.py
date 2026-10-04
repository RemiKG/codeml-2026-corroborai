"""Bounded read-only workbook readers, normalization and safe exports."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel


class InputError(ValueError):
    """An input cannot be safely interpreted."""


def plain(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def ascii_text(value):
    return "".join(c for c in unicodedata.normalize("NFKD", text(value)) if not unicodedata.combining(c))


def key(value):
    return text(value)


def boolean(value):
    if isinstance(value, bool):
        return value
    val = text(value).lower()
    if val in ("1", "true", "oui"):
        return True
    if val in ("0", "false", "non"):
        return False
    raise InputError("Expected a boolean (true/false, Oui/Non or 1/0).")


def number(value):
    if value is None or text(value) == "":
        return None
    try:
        val = float(text(value))
    except ValueError as exc:
        raise InputError("Expected a finite number.") from exc
    if not math.isfinite(val):
        raise InputError("Expected a finite number.")
    return val


def day(value, epoch):
    if value is None or text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    val = text(value)
    if re.fullmatch(r"\d+(?:\.\d+)?", val):
        serial = float(val)
        if not 1 <= serial <= 2958465 or not serial.is_integer():
            raise InputError("Expected an integral, valid Excel date serial.")
        decoded = from_excel(serial, epoch)
        if not isinstance(decoded, datetime):
            raise InputError("Excel time without a date is not supported.")
        return decoded.date()
    try:
        decoded = datetime.fromisoformat(val)
        return decoded.date()
    except ValueError:
        try:
            return date.fromisoformat(val)
        except ValueError as exc:
            raise InputError("Expected ISO date, typed Excel date or Excel serial; ambiguous formats require review.") from exc


def workbook(raw: bytes, name: str):
    if len(raw) > 20_000_000:
        raise InputError(f"{name}: workbook exceeds 20 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100_000_000:
                raise InputError(f"{name}: expanded workbook exceeds 100 MB.")
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except InputError:
        raise
    except Exception as exc:
        raise InputError(f"{name}: cannot read XLSX workbook.") from exc


def sheet_rows(wb, sheet, name):
    if sheet not in wb.sheetnames:
        raise InputError(f"{name}: missing sheet {sheet}.")
    ws = wb[sheet]
    if ws.max_row > 50_000 or ws.max_column > 400:
        raise InputError(f"{name}: worksheet dimensions exceed supported limits.")
    result = []
    for row in ws.iter_rows():
        if any(c.data_type == "f" for c in row):
            raise InputError(f"{name}/{sheet}: formulas are not accepted as cached input values.")
        result.append([c.value for c in row])
    if not result:
        raise InputError(f"{name}/{sheet}: empty worksheet.")
    return result


def table(raw, name, sheet, required, csv_in_cell=False):
    wb = workbook(raw, name)
    try:
        raw_rows = sheet_rows(wb, sheet, name)
        if csv_in_cell and isinstance(raw_rows[0][0], str) and "," in raw_rows[0][0]:
            parsed = []
            for n, row in enumerate(raw_rows, 1):
                if not any(v is not None for v in row):
                    continue
                if any(v is not None for v in row[1:]):
                    raise InputError(f"{name}/{sheet} row {n}: CSV-in-cell must occupy column A only.")
                try:
                    decoded = next(csv.reader([str(row[0])], strict=True))
                except csv.Error as exc:
                    raise InputError(f"{name}/{sheet} row {n}: malformed CSV.") from exc
                parsed.append((n, decoded))
        else:
            parsed = [(n, row) for n, row in enumerate(raw_rows, 1) if any(v is not None for v in row)]
        if not parsed or not any(text(h) for h in parsed[0][1]):
            raise InputError(f"{name}/{sheet}: empty worksheet or missing header row; provide the required column headers and assignment/reference data.")
        headers = [text(h) for h in parsed[0][1]]
        while headers and not headers[-1]:
            headers.pop()
        if len(headers) != len(set(headers)) or "" in headers:
            raise InputError(f"{name}/{sheet}: empty or duplicate headers.")
        missing = set(required) - set(headers)
        if missing:
            raise InputError(f"{name}/{sheet}: missing columns: {', '.join(sorted(missing))}.")
        rows = []
        for rownum, values in parsed[1:]:
            if len(values) > len(headers) and any(v is not None for v in values[len(headers):]):
                raise InputError(f"{name}/{sheet} row {rownum}: unexpected extra fields.")
            if csv_in_cell and len(values) != len(headers):
                raise InputError(f"{name}/{sheet} row {rownum}: expected {len(headers)} CSV fields.")
            item = dict(zip(headers, values[:len(headers)]))
            item.update(_row=rownum, _file=name, _sheet=sheet, _epoch=wb.epoch)
            rows.append(item)
        return rows, {"name": name, "sheet": sheet, "rows": len(rows), "headers": headers,
                      "sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "epoch": wb.epoch.isoformat()}
    finally:
        wb.close()


def ref(row, fields=None):
    if row is None:
        return None
    return {"file": row["_file"], "sheet": row["_sheet"], "row": row["_row"], "fields": list(fields or [])}


EXPORT_COLUMNS = ["id", "person_id", "assignment", "field", "verdict", "source_value", "normalized_source",
                  "expected", "target_value", "normalized_target", "rule_id", "reason", "match_method",
                  "rule_ref", "source_ref", "target_ref", "supporting_refs", "analysis_origin", "suggestions"]


def serialized(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=plain)
    return "" if value is None else str(plain(value))


def safe_csv(value):
    val = serialized(value)
    if val.lstrip().startswith(("=", "+", "-", "@")) or val.startswith(("\t", "\r", "\n")):
        return "'" + val
    return val


def csv_bytes(run):
    out = io.StringIO(newline="")
    writer = csv.writer(out)
    writer.writerow(EXPORT_COLUMNS)
    for result in run["results"]:
        writer.writerow([safe_csv(result.get(c)) for c in EXPORT_COLUMNS])
    return out.getvalue().encode("utf-8-sig")


def xlsx_bytes(run):
    wb = Workbook()
    wb.remove(wb.active)
    groups = [("Investigation", [r for r in run["results"] if r["verdict"] in ("anomaly", "review_required")]),
              ("Justified", [r for r in run["results"] if r["verdict"] == "justified_difference"]),
              ("Audit", run["results"])]
    for title, results in groups:
        ws = wb.create_sheet(title)
        ws.append(EXPORT_COLUMNS)
        for result in results:
            ws.append([serialized(result.get(c)) for c in EXPORT_COLUMNS])
            for cell in ws[ws.max_row]:
                cell.data_type = "s"  # Never interpret an input string as an Excel formula.
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173B3E")
        for name in ("A", "B", "C", "D", "E", "K"):
            ws.column_dimensions[name].width = 24
        ws.column_dimensions["L"].width = 90
    meta = wb.create_sheet("Run metadata")
    for field, value in run["metadata"].items():
        meta.append([field, serialized(value)])
        for cell in meta[meta.max_row]:
            cell.data_type = "s"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def write_outputs(run, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "corroboration.json").write_text(json.dumps(run, ensure_ascii=False, indent=2, default=plain) + "\n", encoding="utf-8")
    (folder / "corroboration.csv").write_bytes(csv_bytes(run))
    (folder / "corroboration.xlsx").write_bytes(xlsx_bytes(run))
